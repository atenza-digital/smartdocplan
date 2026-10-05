from __future__ import annotations

from copy import deepcopy
from pathlib import Path

from docx import Document


ROOT = Path(r"C:\Projetos\Atenza\smartdocplan")
TEMPLATE = ROOT / "tmp" / "atenza_doc" / "template_relatorio.docx"
OUTPUT_DIR = ROOT / "docs" / "cliente"
OUTPUT_DOCX = OUTPUT_DIR / "Relatorio_Acompanhamento_SmartDocPlan_Melhorias_Atenza_v1.0_2026-07-14.docx"

FONT_MAP = {
    "AtenzaNeuePowerUltra": "Neue Power Ultra",
    "AtenzaNeuePowerBold": "Neue Power Bold",
    "AtenzaNeuePower": "Neue Power",
    "AtenzaNorticaBold": "Nortica Typeface Bold",
    "AtenzaNortica": "Nortica Typeface",
}


PROJECT_SHEET = {
    (0, 0): ("Projeto", "SmartDocPlan"),
    (0, 1): ("Cliente final", "SmartDocPlan"),
    (0, 2): ("Responsável técnico Atenza", "Ewerton Gomes Almeida"),
    (1, 0): ("Responsável pelo cliente / ponto focal", "Equipe SmartDocPlan"),
    (1, 1): ("Versão do documento", "v1.0"),
    (1, 2): ("Data", "14/07/2026"),
    (2, 0): ("Status geral", "Em análise complementar"),
    (2, 1): ("Ambiente", "Desenvolvimento / Homologação"),
    (2, 2): ("Próxima ação", "Validar prioridades com a SmartDocPlan na rodada de 21/07/2026."),
    (3, 0): ("Decisão aguardada", "Confirmação dos itens que entram na versão final."),
    (3, 1): ("Entregue", "Levantamento interno de melhorias funcionais, de fluxo e de experiência."),
    (3, 2): ("Pendente", "Validação conjunta com o cliente e definição da ordem de execução final."),
    (4, 0): ("Risco", "Melhorias de maior valor podem entrar tarde se a priorização final demorar."),
    (4, 1): ("Depende de aprovação", "Sim, para fechamento da próxima rodada evolutiva."),
}

EXECUTIVE_CARDS = [
    "Em análise complementar / preparação para validação final",
    "Atenza e SmartDocPlan devem revisar as sugestões deste material e confirmar o recorte da versão final na reunião de 21/07/2026.",
    "Aprovação das prioridades que serão incorporadas à rodada final de evolução do sistema.",
]

CURRENT_STATUS = [
    (
        "Jornada de solicitação de RH",
        "Parcialmente aderente",
        "O fluxo principal já orienta a abertura das solicitações, mas ainda pode ficar mais guiado, previsível e confortável para usuários não técnicos.",
    ),
    (
        "Governança documental",
        "Parcialmente aderente",
        "A base de documentos por empresa, colaborador e solicitação já existe, com oportunidade para ficar mais centralizada e mais fácil de acompanhar.",
    ),
    (
        "Conexão entre módulos",
        "Em evolução",
        "Checklist Docs, Matriz Legal, documentos da empresa e requisitos por função podem aparecer de forma mais explícita dentro da mesma jornada operacional.",
    ),
    (
        "Painéis e visão gerencial",
        "Em evolução",
        "Os painéis já apoiam leitura inicial, mas ainda podem evoluir para indicadores mais acionáveis, com filtros e abertura direta das listas relacionadas.",
    ),
    (
        "Experiência por perfil",
        "Parcialmente aderente",
        "Os perfis já estão definidos, porém a navegação ainda pode ficar mais especializada para solicitante, analista, gestor e administrador.",
    ),
    (
        "RH, SST e temas psicossociais",
        "Em evolução",
        "A plataforma já comporta esse escopo e tem espaço para ampliar visibilidade, contexto e apoio operacional ao acompanhamento contínuo.",
    ),
]

SUGGESTIONS = [
    "Reforçar a jornada guiada da solicitação, deixando claro o que já foi concluído, o que ainda falta e por qual motivo cada etapa existe.",
    "Tornar mais evidente o vínculo entre processo, empresa, função, frente/local, checklist documental e requisitos legais, para que o usuário perceba um fluxo único.",
    "Consolidar um dossiê do colaborador com visão central de documentos pessoais, treinamentos, exames, históricos, férias, afastamentos e pendências.",
    "Dar mais protagonismo aos documentos da empresa, como Cartão CNPJ, Contrato Social, PCMSO, PGR, LTCAT e CNO, com status visual, vencimentos e impacto operacional.",
    "Evoluir a avaliação administrativa para permitir registrar com mais naturalidade número, validade, observações e evidências dos itens analisados, especialmente ASO, exames e treinamentos.",
    "Diferenciar de forma mais clara o que é item obrigatório, opcional, pendente ou recomendado, reduzindo dúvidas na abertura e na tratativa das solicitações.",
    "Ampliar alertas e notificações para vencimentos, correções, rejeições, férias próximas, ausências documentais e pendências críticas por empresa.",
    "Tornar os dashboards mais orientados à ação, permitindo clicar nos indicadores, abrir listas filtradas e acompanhar rapidamente o que exige atenção.",
    "Especializar melhor a experiência por perfil de acesso, simplificando a jornada do solicitante e fortalecendo a visão de gestão para administradores e analistas.",
    "Reforçar o apoio operacional a temas de RH, SST e fatores psicossociais, em linha com o papel da plataforma como organizadora de evidências, requisitos e acompanhamento contínuo.",
]

NEXT_STEPS = [
    (
        "Validar com a SmartDocPlan as sugestões consolidadas neste documento.",
        "Atenza + SmartDocPlan",
        "Realizar leitura conjunta e recolher complementos da equipe cliente.",
        "21/07/2026",
    ),
    (
        "Classificar os itens em essencial, alta prioridade e evolutivo.",
        "Atenza + SmartDocPlan",
        "Fechar o recorte da versão final e o que fica como evolução posterior.",
        "Após reunião",
    ),
    (
        "Confirmar regras finais de documentos por processo, empresa e perfil.",
        "SmartDocPlan",
        "Formalizar parâmetros operacionais para reduzir dúvidas no uso diário.",
        "A definir",
    ),
    (
        "Validar critérios de alertas, vencimentos e pendências críticas.",
        "Atenza + SmartDocPlan",
        "Traduzir regras de acompanhamento para uma experiência mais proativa.",
        "A definir",
    ),
    (
        "Aprovar a ordem de execução da rodada final de melhorias.",
        "SmartDocPlan",
        "Autorizar a sequência de implementação e homologação final.",
        "A definir",
    ),
]

VERSION_HISTORY = [
    ("v1.0", "14/07/2026", "Ewerton Gomes Almeida", "Emissão inicial do relatório com recomendações Atenza para a evolução funcional do SmartDocPlan."),
]


def set_paragraph_text(paragraph, text: str) -> None:
    if paragraph.runs:
        paragraph.runs[0].text = text
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(text)


def set_key_value_cell(cell, label: str, value: str) -> None:
    paragraph = cell.paragraphs[0]
    while len(paragraph.runs) < 4:
        paragraph.add_run("")
    paragraph.runs[0].text = ""
    paragraph.runs[1].text = label
    paragraph.runs[2].text = "\n"
    paragraph.runs[3].text = value
    for run in paragraph.runs[4:]:
        run.text = ""


def set_single_cell_text(cell, text: str) -> None:
    paragraph = cell.paragraphs[0]
    if not paragraph.runs:
        paragraph.add_run(text)
        return
    target = paragraph.runs[-1]
    target.text = text
    for run in paragraph.runs[:-1]:
        if run.text != "\n":
            run.text = ""


def append_cloned_row(table, source_row_index: int):
    table._tbl.append(deepcopy(table.rows[source_row_index]._tr))
    return table.rows[-1]


def fill_table_rows(table, rows, source_row_index: int, start_row_index: int = 1) -> None:
    existing_slots = len(table.rows) - start_row_index
    if existing_slots < len(rows):
        for _ in range(len(rows) - existing_slots):
            append_cloned_row(table, source_row_index)

    for row_index, values in enumerate(rows, start=start_row_index):
        row = table.rows[row_index]
        for cell, value in zip(row.cells, values):
            set_single_cell_text(cell, value)


def normalize_font_name(run) -> None:
    current_name = run.font.name
    if current_name in FONT_MAP:
        run.font.name = FONT_MAP[current_name]


def normalize_fonts(doc: Document) -> None:
    for paragraph in doc.paragraphs:
        for run in paragraph.runs:
            normalize_font_name(run)

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        normalize_font_name(run)

    for section in doc.sections:
        for paragraph in section.header.paragraphs:
            for run in paragraph.runs:
                normalize_font_name(run)
        for paragraph in section.footer.paragraphs:
            for run in paragraph.runs:
                normalize_font_name(run)


def main() -> None:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    doc = Document(str(TEMPLATE))

    paragraphs = doc.paragraphs
    set_paragraph_text(paragraphs[1], "Relatório de Acompanhamento do Projeto")
    set_paragraph_text(
        paragraphs[2],
        "SmartDocPlan: sugestões Atenza para evolução funcional, experiência de uso e valor percebido na versão final.",
    )
    set_paragraph_text(paragraphs[5], (
        "A Atenza realizou uma análise interna do SmartDocPlan com foco em experiência de uso, clareza operacional, valor percebido e aderência ao propósito da ferramenta. "
        "A plataforma já possui uma base relevante para gestão de solicitações, documentos, requisitos legais, operação multiempresa e acompanhamento de rotinas de RH."
    ))
    set_paragraph_text(paragraphs[6], "2. Situação atual")
    set_paragraph_text(paragraphs[7], "3. Sugestões Atenza para a versão final")
    set_paragraph_text(paragraphs[9], "4. Próximos passos para validação")
    set_paragraph_text(paragraphs[10], "5. Observação de rotina")
    set_paragraph_text(paragraphs[11], (
        "Este material foi preparado pela Atenza como complemento proativo à rodada de validação do projeto. "
        "A intenção é agregar valor à versão final do SmartDocPlan, registrando oportunidades percebidas internamente antes do alinhamento conjunto previsto para 21/07/2026."
    ))
    set_paragraph_text(paragraphs[13], "Histórico de versões")

    project_table = doc.tables[0]
    for (row, col), (label, value) in PROJECT_SHEET.items():
        set_key_value_cell(project_table.rows[row].cells[col], label, value)

    cards_table = doc.tables[1]
    for cell, value in zip(cards_table.rows[0].cells, ["STATUS", "PRÓXIMA AÇÃO", "DECISÃO AGUARDADA"]):
        set_single_cell_text(cell, value)
    for cell, value in zip(cards_table.rows[1].cells, EXECUTIVE_CARDS):
        set_single_cell_text(cell, value)

    status_header = doc.tables[2].rows[0].cells
    for cell, value in zip(status_header, ["Frente", "Status", "Observação"]):
        set_single_cell_text(cell, value)
    fill_table_rows(doc.tables[2], CURRENT_STATUS, source_row_index=1, start_row_index=1)

    suggestion_rows = [(str(index), text) for index, text in enumerate(SUGGESTIONS, start=1)]
    fill_table_rows(doc.tables[3], suggestion_rows, source_row_index=0, start_row_index=0)

    next_steps_header = doc.tables[4].rows[0].cells
    for cell, value in zip(next_steps_header, ["Pendência", "Responsável", "Próxima ação", "Prazo"]):
        set_single_cell_text(cell, value)
    fill_table_rows(doc.tables[4], NEXT_STEPS, source_row_index=1, start_row_index=1)

    history_header = doc.tables[5].rows[0].cells
    for cell, value in zip(history_header, ["Versão", "Data", "Responsável", "Descrição"]):
        set_single_cell_text(cell, value)
    fill_table_rows(doc.tables[5], VERSION_HISTORY, source_row_index=1, start_row_index=1)

    normalize_fonts(doc)
    doc.save(str(OUTPUT_DOCX))
    print(OUTPUT_DOCX)


if __name__ == "__main__":
    main()
