from __future__ import annotations

from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.shared import Inches
from docx.text.paragraph import Paragraph
from PIL import Image, ImageDraw, ImageFont


ROOT = Path(r"C:\Projetos\Atenza\smartdocplan")
TEMPLATE = (
    Path(r"C:\Projetos\Documentações\02_Templates\01_Operacionais")
    / "Template_Relatorio_Acompanhamento_Projeto_Atenza.docx"
)
OUTPUT_DIR = ROOT / "docs" / "cliente"
OUTPUT_DOCX = OUTPUT_DIR / "Alinhamento_Melhorias_Estimativa_SmartDocPlan_v1.0_2026-08-25.docx"
ASSET_DIR = OUTPUT_DIR / "assets_alinhamento_melhorias_smartdocplan_2026-08-25"
FONT_DIR = Path(r"C:\Projetos\Documentações\04_Assets\fonts")

FONT_MAP = {
    "AtenzaNeuePowerUltra": "Neue Power Ultra",
    "AtenzaNeuePowerBold": "Neue Power Bold",
    "AtenzaNeuePower": "Neue Power",
    "AtenzaNorticaBold": "Nortica Typeface Bold",
    "AtenzaNortica": "Nortica Typeface",
}

PROJECT_SHEET = {
    (0, 0): ("Projeto", "SmartDocPlan"),
    (0, 1): ("Cliente final", "SmartDocPlan"),
    (0, 2): ("Responsável Atenza", "Ewerton Gomes Almeida"),
    (1, 0): ("Responsável pelo cliente / ponto focal", "Equipe SmartDocPlan"),
    (1, 1): ("Versão do documento", "v1.0"),
    (1, 2): ("Data", "25/08/2026"),
    (2, 0): ("Status geral", "Em homologação - ajustes em análise"),
    (2, 1): ("Ambiente", "Homologação"),
    (2, 2): ("Próxima ação", "Validar entendimento e premissas"),
    (3, 0): ("Decisão aguardada", "Confirmação das regras de negócio críticas"),
    (3, 1): ("Entregue", "Análise das melhorias solicitadas"),
    (3, 2): ("Pendente", "Validação do cliente e priorização"),
    (4, 0): ("Risco", "Decisões pendentes podem alterar a previsão"),
    (4, 1): ("Depende de aprovação", "Sim, para início das etapas correspondentes"),
}

EXECUTIVE_CARDS = [
    "Em homologação. O fluxo atual será aprimorado sem reposicioná-lo como um novo projeto.",
    "Validar com a SmartDocPlan as regras que influenciam a ordem de execução e o prazo das melhorias.",
    "Confirmar os cinco temas listados neste documento para manter a previsão global de 16 semanas.",
]

CURRENT_STATE_ROWS = [
    (
        "Fluxo atual de solicitações",
        "Parcialmente aderente",
        "Já existe base de solicitações, documentos e dossiê; o ajuste é ampliar as regras automáticas e o histórico operacional.",
    ),
    (
        "Matriz legal e requisitos",
        "Parcialmente aderente",
        "Já há requisitos por função; a evolução passa a considerar também processo e frente/local, com origem explicável.",
    ),
    (
        "Férias",
        "Evolução prevista",
        "Será incluído como módulo de gestão com calendário, programação, alertas, filtros e histórico.",
    ),
    (
        "Atestados e afastamentos",
        "Evolução prevista",
        "Será incluído como módulo com controle de retorno, anexos, histórico e permissões adequadas aos dados sensíveis.",
    ),
    (
        "Notificações, BI e relatórios",
        "Parcialmente aderente",
        "A central de notificações já existe; os indicadores, filtros e detalhamentos serão ampliados com as novas rotinas.",
    ),
]

DELIVERY_ROWS = [
    (
        "1",
        "Aprimorar frentes/locais, histórico de função e local do colaborador e motor de requisitos por processo, função e frente/local.",
    ),
    (
        "2",
        "Aprimorar solicitações e dossiê: requisitos automáticos, origem explicável, fotografia da regra aplicada, pendências, documentos, histórico e conformidade transparente.",
    ),
    (
        "3",
        "Incluir gestão de férias: períodos aquisitivo e concessivo, programação, aprovação, calendário, alertas, filtros e histórico por colaborador.",
    ),
    (
        "4",
        "Incluir atestados e afastamentos: registro, anexos, período, retorno previsto, pendências, histórico e permissões reforçadas para dados sensíveis.",
    ),
    (
        "5",
        "Ampliar BI, alertas e relatórios com indicadores rastreáveis e filtros por empresa, função, frente/local, período e status.",
    ),
]

DECISION_ROWS = [
    (
        "Frente/local de trabalho",
        "SmartDocPlan",
        "Confirmar se frente, obra, unidade e contrato são uma única entidade ou cadastros diferentes.",
        "Até 5 dias úteis após o envio",
    ),
    (
        "Matriz de requisitos",
        "SmartDocPlan / Atenza",
        "Definir se a regra será por processo, função, frente/local ou combinação desses critérios.",
        "Até 5 dias úteis após o envio",
    ),
    (
        "Conformidade do dossiê",
        "SmartDocPlan / Atenza",
        "Aprovar percentual, pesos e tratamento de itens pendentes ou vencidos.",
        "Até o início da Etapa 2",
    ),
    (
        "Férias",
        "SmartDocPlan",
        "Definir aprovação e se haverá cálculo de saldo legal ou apenas registro de dados validados.",
        "Até o início da Etapa 3",
    ),
    (
        "Dados de saúde",
        "SmartDocPlan / Atenza",
        "Definir dados permitidos e perfis autorizados a visualizar atestados e afastamentos.",
        "Até o início da Etapa 4",
    ),
]

VERSION_HISTORY = [
    (
        "v1.0",
        "25/08/2026",
        "Ewerton Gomes Almeida",
        "Emissão inicial para confirmação das melhorias do fluxo atual em homologação e da estimativa por etapas.",
    ),
]


FLOW_DIAGRAMS = [
    (
        "Fluxo 1 - Jornada de solicitação e conformidade",
        "Os requisitos passam a ser definidos por regras da empresa e acompanhados até o dossiê do colaborador.",
        [
            ("Configuração", "Funções, frentes e matriz"),
            ("Solicitação", "Processo, pessoa e contexto"),
            ("Requisitos", "Documentos, exames e treinamentos"),
            ("Análise", "Envio, validação e correção"),
            ("Dossiê", "Histórico, pendências e conformidade"),
        ],
        "fluxo_solicitacao_conformidade.png",
    ),
    (
        "Fluxo 2 - Rotinas trabalhistas integradas",
        "Férias, atestados e afastamentos passam a alimentar os alertas, o histórico e os indicadores da plataforma.",
        [
            ("Colaborador", "Vínculo, função e frente/local"),
            ("Férias", "Programação, aprovação e calendário"),
            ("Atestados", "Registro, retorno e acesso protegido"),
            ("Alertas e BI", "Pendências, notificações e visão gerencial"),
        ],
        "fluxo_rotinas_trabalhistas.png",
    ),
]


def set_paragraph_text(paragraph, text: str) -> None:
    if paragraph.runs:
        paragraph.runs[0].text = text
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(text)


def set_single_cell_text(cell, text: str) -> None:
    paragraph = cell.paragraphs[0]
    if not paragraph.runs:
        paragraph.add_run(text)
        return
    paragraph.runs[-1].text = text
    for run in paragraph.runs[:-1]:
        if run.text != "\n":
            run.text = ""


def set_key_value_cell(cell, label: str, value: str) -> None:
    paragraph = cell.paragraphs[0]
    while len(paragraph.runs) < 4:
        paragraph.add_run("")
    paragraph.runs[0].text = ""
    paragraph.runs[1].text = label
    paragraph.runs[2].text = "\n"
    paragraph.runs[3].text = value
    for run in paragraph.runs[4:]:
        run.text = ""


def append_cloned_row(table, source_row_index: int):
    table._tbl.append(deepcopy(table.rows[source_row_index]._tr))
    return table.rows[-1]


def insert_paragraph_after(anchor: Paragraph) -> Paragraph:
    new_p = OxmlElement("w:p")
    anchor._p.addnext(new_p)
    return Paragraph(new_p, anchor._parent)


def clone_paragraph_after(anchor: Paragraph, source: Paragraph, text: str) -> Paragraph:
    new_p = deepcopy(source._p)
    anchor._p.addnext(new_p)
    paragraph = Paragraph(new_p, anchor._parent)
    set_paragraph_text(paragraph, text)
    return paragraph


def fill_table_rows(table, rows, source_row_index: int, start_row_index: int = 1) -> None:
    existing_slots = len(table.rows) - start_row_index
    if existing_slots < len(rows):
        for _ in range(len(rows) - existing_slots):
            append_cloned_row(table, source_row_index)

    for row_index, values in enumerate(rows, start=start_row_index):
        row = table.rows[row_index]
        for cell, value in zip(row.cells, values):
            set_single_cell_text(cell, value)


def normalize_fonts(doc: Document) -> None:
    def normalize_paragraphs(paragraphs) -> None:
        for paragraph in paragraphs:
            for run in paragraph.runs:
                if run.font.name in FONT_MAP:
                    run.font.name = FONT_MAP[run.font.name]

    normalize_paragraphs(doc.paragraphs)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                normalize_paragraphs(cell.paragraphs)
    for section in doc.sections:
        normalize_paragraphs(section.header.paragraphs)
        normalize_paragraphs(section.footer.paragraphs)


def wrap_text(draw: ImageDraw.ImageDraw, text: str, font, max_width: int) -> list[str]:
    words = text.split()
    lines: list[str] = []
    line = ""
    for word in words:
        candidate = f"{line} {word}".strip()
        if draw.textbbox((0, 0), candidate, font=font)[2] <= max_width:
            line = candidate
        else:
            if line:
                lines.append(line)
            line = word
    if line:
        lines.append(line)
    return lines


def draw_centered_lines(draw: ImageDraw.ImageDraw, lines: list[str], font, box: tuple[int, int, int, int], color: str, line_gap: int = 8) -> None:
    x1, y1, x2, y2 = box
    heights = [draw.textbbox((0, 0), line, font=font)[3] for line in lines]
    total_height = sum(heights) + line_gap * max(0, len(lines) - 1)
    current_y = y1 + (y2 - y1 - total_height) // 2
    for line, height in zip(lines, heights):
        width = draw.textbbox((0, 0), line, font=font)[2]
        draw.text((x1 + (x2 - x1 - width) // 2, current_y), line, font=font, fill=color)
        current_y += height + line_gap


def create_flow_diagram(title: str, subtitle: str, steps: list[tuple[str, str]], filename: str) -> Path:
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    image_path = ASSET_DIR / filename
    width, height = 1800, 520
    image = Image.new("RGB", (width, height), "#F2F2ED")
    draw = ImageDraw.Draw(image)
    title_font = ImageFont.truetype(str(FONT_DIR / "AtenzaNeuePowerBold.ttf"), 34)
    subtitle_font = ImageFont.truetype(str(FONT_DIR / "AtenzaNortica.ttf"), 20)
    card_title_font = ImageFont.truetype(str(FONT_DIR / "AtenzaNorticaBold.ttf"), 23)
    card_body_font = ImageFont.truetype(str(FONT_DIR / "AtenzaNortica.ttf"), 18)

    draw.text((66, 48), title, font=title_font, fill="#030409")
    draw.text((66, 98), subtitle, font=subtitle_font, fill="#3850A0")
    draw.line((66, 142, width - 66, 142), fill="#00FFCC", width=4)

    count = len(steps)
    margin = 66
    arrow_width = 42
    gap = 28
    available_width = width - (margin * 2) - (arrow_width * (count - 1)) - (gap * (count - 1))
    card_width = available_width // count
    card_height = 220
    top = 220

    for index, (label, description) in enumerate(steps):
        left = margin + index * (card_width + gap + arrow_width)
        right = left + card_width
        draw.rounded_rectangle((left, top, right, top + card_height), radius=18, fill="#FFFFFF", outline="#030409", width=2)
        draw.rectangle((left, top, right, top + 12), fill="#00FFCC")
        draw_centered_lines(draw, wrap_text(draw, label, card_title_font, card_width - 34), card_title_font, (left + 18, top + 40, right - 18, top + 112), "#030409")
        draw_centered_lines(draw, wrap_text(draw, description, card_body_font, card_width - 36), card_body_font, (left + 18, top + 112, right - 18, top + card_height - 24), "#3850A0", 7)

        if index < count - 1:
            arrow_left = right + gap
            arrow_top = top + card_height // 2
            draw.line((arrow_left, arrow_top, arrow_left + arrow_width - 10, arrow_top), fill="#3850A0", width=7)
            draw.polygon(
                [
                    (arrow_left + arrow_width - 10, arrow_top - 15),
                    (arrow_left + arrow_width + 8, arrow_top),
                    (arrow_left + arrow_width - 10, arrow_top + 15),
                ],
                fill="#3850A0",
            )

    image.save(image_path, quality=95)
    return image_path


def insert_flow_diagrams(heading: Paragraph) -> Paragraph:
    anchor = heading
    for title, subtitle, steps, filename in FLOW_DIAGRAMS:
        image_path = create_flow_diagram(title, subtitle, steps, filename)
        image_paragraph = insert_paragraph_after(anchor)
        image_paragraph.alignment = 1
        image_paragraph.add_run().add_picture(str(image_path), width=Inches(6.35))
        anchor = image_paragraph
    return anchor


def main() -> None:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    doc = Document(str(TEMPLATE))
    paragraphs = doc.paragraphs

    set_paragraph_text(paragraphs[1], "Alinhamento de Melhorias e Estimativa")
    set_paragraph_text(
        paragraphs[2],
        "SmartDocPlan: confirmação do entendimento das melhorias do fluxo atual em homologação e previsão de execução por etapas.",
    )
    set_paragraph_text(paragraphs[4], "1. Resumo executivo")
    set_paragraph_text(
        paragraphs[5],
        "A Atenza consolidou o entendimento das melhorias apresentadas para a versão atual do SmartDocPlan, que permanece em homologação. O objetivo é aprimorar o fluxo já existente de solicitações, documentos e dossiê, além de incluir as rotinas de férias e de atestados/afastamentos como evoluções integradas à plataforma.",
    )
    set_paragraph_text(paragraphs[6], "2. Situação atual e entendimento")
    set_paragraph_text(paragraphs[7], "3. Fluxos de evolução")
    set_paragraph_text(paragraphs[9], "5. Confirmações necessárias e próximos passos")
    paragraphs[9].paragraph_format.page_break_before = True
    paragraphs[9].paragraph_format.keep_with_next = True
    set_paragraph_text(paragraphs[10], "6. Previsão de prazo")
    paragraphs[10].paragraph_format.page_break_before = True
    paragraphs[10].paragraph_format.keep_with_next = True
    set_paragraph_text(
        paragraphs[11],
        "A previsão para as melhorias mapeadas é de 16 semanas, contadas após a confirmação do entendimento e das decisões críticas. A execução será organizada em: Etapa 1 - regras e vínculos (3 semanas); Etapa 2 - solicitações, dossiê e conformidade (3 semanas); Etapa 3 - férias (4 semanas); Etapa 4 - atestados e afastamentos (3 semanas); e Etapa 5 - BI, alertas e relatórios (3 semanas). Integrações externas, cálculos financeiros de férias ou novas demandas serão avaliados separadamente caso sejam solicitados.\n\nSolicita-se a confirmação de que o entendimento acima representa a necessidade da SmartDocPlan e o retorno consolidado sobre as decisões listadas. Com essa validação, a Atenza organizará o backlog executivo e seguirá com a primeira etapa priorizada para a versão em homologação.",
    )
    set_paragraph_text(paragraphs[13], "Histórico de versões")

    project_table = doc.tables[0]
    for (row, col), (label, value) in PROJECT_SHEET.items():
        set_key_value_cell(project_table.rows[row].cells[col], label, value)

    cards_table = doc.tables[1]
    for cell, value in zip(cards_table.rows[1].cells, EXECUTIVE_CARDS):
        set_single_cell_text(cell, value)

    current_state_table = doc.tables[2]
    fill_table_rows(current_state_table, CURRENT_STATE_ROWS, source_row_index=1, start_row_index=1)

    delivery_table = doc.tables[3]
    fill_table_rows(delivery_table, DELIVERY_ROWS, source_row_index=0, start_row_index=0)

    last_flow = insert_flow_diagrams(paragraphs[7])
    improvements_heading = clone_paragraph_after(last_flow, paragraphs[7], "4. Melhorias priorizadas")
    improvements_heading.paragraph_format.keep_with_next = True

    decisions_table = doc.tables[4]
    for cell, value in zip(decisions_table.rows[0].cells, ("Pendência", "Responsável", "Próxima ação", "Prazo")):
        set_single_cell_text(cell, value)
    fill_table_rows(decisions_table, DECISION_ROWS, source_row_index=1, start_row_index=1)

    history_table = doc.tables[5]
    for cell, value in zip(history_table.rows[0].cells, ("Versão", "Data", "Responsável", "Descrição")):
        set_single_cell_text(cell, value)
    fill_table_rows(history_table, VERSION_HISTORY, source_row_index=1, start_row_index=1)

    for cell, value in zip(current_state_table.rows[0].cells, ("Frente", "Status", "Entendimento Atenza")):
        set_single_cell_text(cell, value)

    for row_index, values in enumerate(CURRENT_STATE_ROWS, start=1):
        for cell, value in zip(current_state_table.rows[row_index].cells, values):
            set_single_cell_text(cell, value)

    normalize_fonts(doc)
    doc.save(str(OUTPUT_DOCX))
    print(OUTPUT_DOCX)


if __name__ == "__main__":
    main()
