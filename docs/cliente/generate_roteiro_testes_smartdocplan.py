from __future__ import annotations

from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.shared import Inches
from docx.text.paragraph import Paragraph


ROOT = Path(r"C:\Projetos\Atenza\smartdocplan")
TEMPLATE = ROOT / "tmp" / "atenza_doc" / "template_roteiro_testes.docx"
OUTPUT_DIR = ROOT / "docs" / "cliente"
OUTPUT_DOCX = OUTPUT_DIR / "Roteiro_Testes_Validacao_SmartDocPlan_v1.0_2026-07-30.docx"
SCREENSHOT_DIR = OUTPUT_DIR / "roteiro_testes_smartdocplan_2026-07-30"

FONT_MAP = {
    "AtenzaNeuePowerUltra": "Neue Power Ultra",
    "AtenzaNeuePowerBold": "Neue Power Bold",
    "AtenzaNeuePower": "Neue Power",
    "AtenzaNorticaBold": "Nortica Typeface Bold",
    "AtenzaNortica": "Nortica Typeface",
}

PROJECT_SHEET = {
    (0, 0): ("Projeto", "SmartDocPlan"),
    (0, 1): ("Cliente final", "SmartDocPlan"),
    (0, 2): ("Responsável Atenza", "Ewerton Gomes Almeida"),
    (1, 0): ("Responsável pela homologação", "Equipe SmartDocPlan / Atenza"),
    (1, 1): ("Ambiente testado", "Homologação publicada na VPS"),
    (1, 2): ("Versão/build", "6b80c5f - deploy CI/CD 30/07/2026"),
    (2, 0): ("Data da rodada", "30/07/2026"),
    (2, 1): ("Prazo de retorno", "01/08/2026"),
    (2, 2): ("Status da validação", "Em homologação"),
    (3, 0): ("Critério de aprovação", "Cenários críticos executados sem bloqueio funcional"),
    (3, 1): ("Decisão aguardada", "Validação do cliente sobre a rodada atual"),
}

EXECUTIVE_CARDS = [
    "Em homologação da rodada publicada em 30/07/2026.",
    "Retorno esperado até 01/08/2026 com apontamentos funcionais, bloqueios ou aceite da rodada.",
    "Aprovação dos cenários CT-001 a CT-007 sem erro crítico, regressão ou falha impeditiva.",
]

ENVIRONMENT_ROWS = [
    ("URL de homologação", "http://89.116.214.65:8080"),
    ("URL de produção, se aplicável", "Não aplicável nesta rodada"),
    ("Usuário/perfil de teste", "platform_admin - admin@smartdocplan.com"),
    ("Senha", "Enviada por canal seguro"),
    ("Versão/build", "Commit 6b80c5f / publicação via GitHub Actions"),
]

SCOPE_ROWS = [
    (
        "1",
        "Fluxos que serão testados: login administrativo, dashboard da plataforma, troca para visão da empresa, configurações da empresa, notificações e abertura da nova solicitação.",
    ),
    (
        "2",
        "Itens fora desta rodada: recuperação de senha, fluxo público seguro, score de conformidade completo, notificações por e-mail, retenção/descarte documental e governança avançada de sessão.",
    ),
    (
        "3",
        "Dados de teste necessários: ambiente VPS acessível, navegador desktop atualizado, perfil platform_admin e empresa de teste já cadastrada.",
    ),
    (
        "4",
        "Critério para aprovação ou reprovação: as telas devem carregar sem erro interno, as navegações devem manter a sessão ativa e os fluxos descritos precisam ficar utilizáveis para continuação da homologação.",
    ),
]

SCENARIOS = [
    (
        "CT-001",
        "Login do administrador da plataforma",
        "Acessar /login, informar e-mail do platform_admin, informar a senha recebida por canal seguro e clicar em Entrar.",
        "O sistema deve autenticar com sucesso e redirecionar para o Dashboard Global sem exibir erro interno.",
        "Pendente",
    ),
    (
        "CT-002",
        "Dashboard Global da plataforma",
        "Após o login, conferir os cards principais, menu lateral, empresa em foco no topo e acesso às áreas administrativas.",
        "O dashboard deve abrir com menu funcional, indicadores visíveis e navegação estável.",
        "Pendente",
    ),
    (
        "CT-003",
        "Entrada na visão da empresa",
        "A partir da plataforma, acessar a visão da empresa pelo seletor superior e entrar na empresa de teste.",
        "A aplicação deve trocar para o contexto da empresa, exibindo o painel de RH e o botão de retorno para a plataforma.",
        "Pendente",
    ),
    (
        "CT-004",
        "Configurações e atualização cadastral da empresa",
        "Na visão da empresa, abrir Configurações e conferir o histórico de atualização cadastral, dados da empresa e abas de apoio.",
        "A tela deve carregar sem erro e exibir dados da empresa, histórico de solicitação e navegação entre abas.",
        "Pendente",
    ),
    (
        "CT-005",
        "Documentos da empresa e pendências documentais",
        "Na aba Empresa de Configurações, localizar o bloco de documentos da empresa e verificar indicadores de obrigatórios pendentes.",
        "Os documentos corporativos devem aparecer com status visual, ação de anexo e contadores coerentes de pendência.",
        "Pendente",
    ),
    (
        "CT-006",
        "Central de notificações",
        "Na visão da empresa, abrir o ícone de notificações e conferir o card com a atualização cadastral aguardando tratativa.",
        "A central deve listar a notificação existente, sem quebrar a tela nem encerrar a sessão.",
        "Pendente",
    ),
    (
        "CT-007",
        "Nova solicitação em etapas",
        "Abrir Nova solicitação e conferir o fluxo em etapas, resumo lateral e início do preenchimento pela etapa Empresa.",
        "A jornada deve aparecer segmentada em etapas visíveis, com layout utilizável e caminho claro para seguir o fluxo.",
        "Pendente",
    ),
]

EVIDENCES = [
    ("Tela 1 - Login", SCREENSHOT_DIR / "01-login.png"),
    ("Tela 2 - Dashboard Global da plataforma", SCREENSHOT_DIR / "02-dashboard-plataforma.png"),
    ("Tela 3 - Painel da empresa", SCREENSHOT_DIR / "04-dashboard-empresa.png"),
    ("Tela 4 - Configurações da empresa", SCREENSHOT_DIR / "05-configuracoes-empresa.png"),
    ("Tela 5 - Central de notificações", SCREENSHOT_DIR / "06-notificacoes.png"),
    ("Tela 6 - Nova solicitação em etapas", SCREENSHOT_DIR / "07-nova-solicitacao.png"),
]

VERSION_HISTORY = [
    (
        "v1.0",
        "30/07/2026",
        "Ewerton Gomes Almeida",
        "Emissão inicial do roteiro de testes da rodada publicada via CI/CD, com evidências visuais e cenários críticos para homologação.",
    ),
]


def set_paragraph_text(paragraph, text: str) -> None:
    if paragraph.runs:
        paragraph.runs[0].text = text
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(text)


def set_single_cell_text(cell, text: str) -> None:
    paragraph = cell.paragraphs[0]
    if not paragraph.runs:
        paragraph.add_run(text)
        return
    target = paragraph.runs[-1]
    target.text = text
    for run in paragraph.runs[:-1]:
        if run.text != "\n":
            run.text = ""


def set_key_value_cell(cell, label: str, value: str) -> None:
    paragraph = cell.paragraphs[0]
    while len(paragraph.runs) < 4:
        paragraph.add_run("")
    paragraph.runs[0].text = ""
    paragraph.runs[1].text = label
    paragraph.runs[2].text = "\n"
    paragraph.runs[3].text = value
    for run in paragraph.runs[4:]:
        run.text = ""


def append_cloned_row(table, source_row_index: int):
    table._tbl.append(deepcopy(table.rows[source_row_index]._tr))
    return table.rows[-1]


def fill_table_rows(table, rows, source_row_index: int, start_row_index: int = 1) -> None:
    existing_slots = len(table.rows) - start_row_index
    if existing_slots < len(rows):
        for _ in range(len(rows) - existing_slots):
            append_cloned_row(table, source_row_index)

    for row_index, values in enumerate(rows, start=start_row_index):
        row = table.rows[row_index]
        for cell, value in zip(row.cells, values):
            set_single_cell_text(cell, value)


def normalize_font_name(run) -> None:
    current_name = run.font.name
    if current_name in FONT_MAP:
        run.font.name = FONT_MAP[current_name]


def normalize_fonts(doc: Document) -> None:
    for paragraph in doc.paragraphs:
        for run in paragraph.runs:
            normalize_font_name(run)

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        normalize_font_name(run)

    for section in doc.sections:
        for paragraph in section.header.paragraphs:
            for run in paragraph.runs:
                normalize_font_name(run)
        for paragraph in section.footer.paragraphs:
            for run in paragraph.runs:
                normalize_font_name(run)


def insert_paragraph_after(anchor: Paragraph, text: str = "", style=None) -> Paragraph:
    new_p = OxmlElement("w:p")
    anchor._p.addnext(new_p)
    paragraph = Paragraph(new_p, anchor._parent)
    if style is not None:
        paragraph.style = style
    if text:
        paragraph.add_run(text)
    return paragraph


def insert_evidences(doc: Document, anchor_index: int) -> None:
    anchor = doc.paragraphs[anchor_index]
    note_style = anchor.style

    for evidence_index, (title, image_path) in enumerate(EVIDENCES):
        anchor = insert_paragraph_after(anchor)

        caption = insert_paragraph_after(anchor, style=note_style)
        caption.add_run(title)
        anchor = caption

        paragraph = insert_paragraph_after(anchor)
        paragraph.add_run().add_picture(str(image_path), width=Inches(6.1))
        anchor = paragraph

        note = insert_paragraph_after(
            anchor,
            "Print representativo da versão publicada em 30/07/2026 para apoiar a conferência visual deste cenário.",
            style=note_style,
        )
        anchor = note

        if evidence_index < len(EVIDENCES) - 1:
            anchor = insert_paragraph_after(anchor)


def main() -> None:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    doc = Document(str(TEMPLATE))

    paragraphs = doc.paragraphs
    set_paragraph_text(paragraphs[1], "Roteiro de Testes e Validação")
    set_paragraph_text(
        paragraphs[2],
        "SmartDocPlan: rodada de homologação da versão publicada em 30/07/2026, com foco nos fluxos críticos já implementados e testáveis nesta etapa.",
    )
    set_paragraph_text(
        paragraphs[9],
        "Os prints abaixo representam a versão publicada na rodada atual e servem como apoio visual para a conferência dos cenários listados neste roteiro.",
    )

    project_table = doc.tables[0]
    for (row, col), (label, value) in PROJECT_SHEET.items():
        set_key_value_cell(project_table.rows[row].cells[col], label, value)

    cards_table = doc.tables[1]
    for cell, value in zip(cards_table.rows[1].cells, EXECUTIVE_CARDS):
        set_single_cell_text(cell, value)

    environment_table = doc.tables[2]
    for idx, values in enumerate(ENVIRONMENT_ROWS, start=1):
        for cell, value in zip(environment_table.rows[idx].cells, values):
            set_single_cell_text(cell, value)

    scope_table = doc.tables[3]
    for idx, values in enumerate(SCOPE_ROWS):
        for cell, value in zip(scope_table.rows[idx].cells, values):
            set_single_cell_text(cell, value)

    scenarios_table = doc.tables[4]
    fill_table_rows(scenarios_table, SCENARIOS, source_row_index=1, start_row_index=1)

    history_table = doc.tables[5]
    fill_table_rows(history_table, VERSION_HISTORY, source_row_index=1, start_row_index=1)

    insert_evidences(doc, anchor_index=9)

    normalize_fonts(doc)
    doc.save(str(OUTPUT_DOCX))
    print(OUTPUT_DOCX)


if __name__ == "__main__":
    main()
